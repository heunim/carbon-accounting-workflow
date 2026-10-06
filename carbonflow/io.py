from io import BytesIO
from pathlib import Path
from dataclasses import asdict
import csv
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .catalog import read_bytes, records
from .engine import aggregate


def load_activities(source, filename="activities.xlsx"):
    raw = read_bytes(source)
    if filename.lower().endswith(".csv"):
        rows = list(csv.DictReader(raw.decode("utf-8-sig").splitlines()))
    else:
        wb = openpyxl.load_workbook(BytesIO(raw), data_only=True, read_only=True)
        sheet = wb["活动数据"] if "活动数据" in wb else wb.worksheets[0]
        rows = records(sheet); wb.close()
    if not rows:
        raise ValueError("活动数据为空")
    required = {"活动名称", "数量", "单位"}
    if not required.issubset(rows[0]):
        raise ValueError("请提供活动名称、数量、单位三列")
    if len(rows) > 10000:
        raise ValueError("演示版本单批最多10000条，请分批上传")
    return rows


def safe(value):
    """避免把输入文本当作Excel公式。"""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def add_sheet(wb, name, headers, rows):
    s = wb.create_sheet(name)
    s.append(headers)
    for row in rows:
        s.append([safe(v) for v in row])
    s.freeze_panes = "A2"; s.auto_filter.ref = s.dimensions
    for cell in s[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174E40")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    s.row_dimensions[1].height = 30
    for row in s.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F6F3")
            if isinstance(cell.value, (float, int)):
                cell.number_format = "0.########"
    for i, h in enumerate(headers, 1):
        s.column_dimensions[get_column_letter(i)].width = 42 if any(x in str(h) for x in ["来源", "公式", "警告", "原因", "说明"]) else 20
    return s


LABELS = {"activity_id": "记录ID", "name": "活动名称", "quantity": "原数量", "unit": "原单位", "row_number": "原行号",
          "status": "状态", "reason": "待处理原因", "scope": "范围", "category": "活动类型", "substance": "物质",
          "factor_id": "因子ID", "factor_value": "因子值", "factor_unit": "因子单位", "standard_quantity": "标准活动量",
          "emission_kg": "排放量kgCO2e", "gas": "排放表示", "gwp": "GWP", "gwp_key": "GWP键", "gwp_basis": "GWP口径",
          "formula": "计算公式", "source": "来源", "source_url": "来源URL", "source_locator": "来源定位",
          "factor_year": "数据年份", "verification": "核对状态", "warnings": "警告", "candidates": "候选因子ID", "rule_id": "归类规则ID"}


def export_excel(results, project, catalog, audit, scenario=None):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    a = aggregate(results)
    add_sheet(wb, "汇总", ["项目", "数值", "说明"], [
        ["项目名称", project.name, ""], ["核算年度", project.year, ""],
        ["已计算排放小计kgCO2e", a["total_kg"], "包含已确认估算，范围2采用位置法；不是完整企业清单"],
        *[[k, v, "kgCO2e"] for k, v in a["scopes"].items()],
        ["已计算条数", a["calculated_count"], ""], ["待处理条数", a["pending_count"], "未计算不等于零排放"],
        ["已排除条数", a["excluded_count"], "请查看明细原因"], ["运行ID", audit.run_id, ""]])
    keys = list(LABELS)
    add_sheet(wb, "计算明细", list(LABELS.values()), [[r.to_dict().get(k) for k in keys] for r in results])
    add_sheet(wb, "待处理清单", ["记录ID", "原行号", "活动名称", "状态", "原因", "候选因子"],
              [[r.activity_id, r.row_number, r.name, r.status, r.reason, r.candidates] for r in results if r.status not in ["已计算", "已排除"]])
    notes = [["报告边界", project.boundary], ["GWP口径", project.gwp_basis], ["因子文件SHA256", catalog.digest],
             ["覆盖说明", "仅CO₂因子未自动补齐CH₄、N₂O。国外代理、缺省参数和人工确认详见每行警告。"],
             ["情景说明", "购电假设情景不改变位置法，不计入主小计；正式市场法未在本版实现。"],
             ["数据性质", "自带示例活动为虚构；实际上传活动的真实性由使用者核验。"]]
    add_sheet(wb, "假设与口径", ["项目", "说明"], notes)
    if scenario:
        add_sheet(wb, "购电情景", ["项目", "值"], list(scenario.items()))
    factor_ids = {r.factor_id for r in results if r.factor_id}
    fs = [f for f in catalog.factors if f["因子ID"] in factor_ids]
    if fs:
        add_sheet(wb, "使用因子快照", list(fs[0]), [list(f.values()) for f in fs])
    gwp_keys = {r.gwp_key for r in results if r.gwp_key and r.status == "已计算"}
    gs = [g for g in catalog.gwps if g["GWP键"] in gwp_keys]
    if gs:
        add_sheet(wb, "使用GWP快照", list(gs[0]), [list(g.values()) for g in gs])
    stream = BytesIO(); wb.save(stream)
    return stream.getvalue()
