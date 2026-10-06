from copy import deepcopy
from io import BytesIO
import json
import math
import openpyxl
import pytest
from carbonflow.engine import run, aggregate
from carbonflow.units import convert, UnitError
from carbonflow.calculator import calculate, electricity_scenario
from carbonflow.catalog import Catalog, CatalogError
from carbonflow.io import export_excel, load_activities
from carbonflow.models import Project


def test_demo_independent_calculation(demo, project, catalog):
    rs, audit = run(demo, project, catalog)
    # 期望值直接来自独立算式，而非被测函数的中间结果。
    s1 = 1000 * (57.2 * .035611) + (200 * .84 / 1000) * (74.1 * 43) + 2 * 1924
    s2 = 10000 * .5581 + 20 * 1000 * .1564
    s3 = 1000 * .07926 + 2000 * .06449 + 100 * .1913 + 2 * 497.28993
    a = aggregate(rs)
    assert a["total_kg"] == pytest.approx(s1 + s2 + s3)
    assert a["scopes"] == pytest.approx({"范围1": s1, "范围2": s2, "范围3": s3})
    assert (a["calculated_count"], a["pending_count"]) == (9, 3)
    assert rs[9].status == "待确认"
    assert rs[10].status == "待补因子"
    assert rs[11].status == "待确认"
    assert sum(a["shares"].values()) == pytest.approx(1)
    assert any("仅CO₂" in w for w in rs[2].warnings)
    assert "factor_snapshot" in audit.jsonl()


@pytest.mark.parametrize("q,a,b,expected", [(2,"t","kg",2000),(2000,"kg","t",2),(2,"MWh","kWh",2000),(5,"GJ","MJ",5000),(1000,"L","m3",1),(1,"m³","L",1000)])
def test_physical_units(q,a,b,expected):
    assert convert(q,a,b)[0] == pytest.approx(expected)


def test_density_only_when_needed():
    assert convert(100,"L","L")[0] == 100
    with pytest.raises(UnitError, match="密度"):
        convert(100,"L","t")
    assert convert(100,"L","t",{"密度kg/L":.84,"密度来源":"测试输入"})[0] == pytest.approx(100*.84/1000)
    with pytest.raises(UnitError, match="来源"):
        convert(100,"L","t",{"密度kg/L":.84})


def test_semantic_units():
    with pytest.raises(UnitError):convert(100,"km","person.km")
    assert convert(100,"km","person.km",{"人数":3,"行程次数":2})[0] == 600
    assert convert(600,"人公里","person.km",{"人数":3})[0] == 600
    with pytest.raises(UnitError):convert(100,"km","vehicle.km")
    with pytest.raises(UnitError):convert(100,"m3","Nm3")
    assert convert(100,"m3","Nm3",{"标准体积已确认":"是"})[0] == 100


@pytest.mark.parametrize("value", [None,"", "abc",float("nan"),float("inf"),-10])
def test_invalid_rows_do_not_block_valid(demo, project, catalog, value):
    rows = [dict(demo[0],数量=value), demo[1]]
    rs, audit = run(rows, project, catalog)
    assert rs[0].status == "无效输入"
    assert rs[1].status == "已计算"
    audit.jsonl()  # 非有限数也不能破坏导出。
    assert export_excel(rs,project,catalog,audit)


def test_missing_factor_not_zero(demo, project, catalog):
    f=catalog.factor("F031-v2023");f["因子值"]=""
    rs,_=run([demo[0],demo[1]],project,catalog)
    assert rs[0].status == "待补因子" and rs[0].emission_kg is None
    assert aggregate(rs)["total_kg"] == pytest.approx(20000*.1564)
    f["因子值"]=0
    rs,_=run([demo[0]],project,catalog)
    assert rs[0].status == "已计算" and rs[0].emission_kg == 0
    assert aggregate(rs)["shares"] == {"范围1":0,"范围2":0,"范围3":0}


def test_gwp_paths(project,catalog):
    co2 = catalog.factor("F031-v2023")
    assert calculate(10,co2,catalog,project)[0] == pytest.approx(10*.5581)
    eq = catalog.factor("UK26-27_304_3136_14_1")
    assert calculate(1000,eq,catalog,project)[0] == pytest.approx(1000*.07926)
    gas = catalog.factor("GAS26-R410A")
    assert calculate(2,gas,catalog,project)[0] == 2*1924
    # 使用实际GWP28检验明确气体质量的折算，不作为生产因子使用。
    ch4 = dict(gas,排放表示="CH4",GWP键="AR5-CH4",气体质量单位="g")
    assert calculate(1000,ch4,catalog,project)[0] == 1000*.001*28
    with pytest.raises(UnitError):calculate(2,gas,catalog,Project(gwp_basis="AR6-100"))
    with pytest.raises(UnitError):calculate(2,eq,catalog,Project(gwp_basis="AR6-100"))


def test_keywords_and_control(project,catalog):
    rows=[{"活动名称":x,"数量":100,"单位":"L","设施归属":"边界内"} for x in ["柴油","煤油","电锅炉用电"]]
    rows[2]["单位"]="kWh"
    rs,_=run(rows,project,catalog)
    assert rs[0].status=="待确认" and rs[1].status=="待确认"
    assert rs[2].scope=="范围2" and rs[2].status=="已计算"
    rows[2]["设施归属"]="边界外"
    assert run(rows,project,catalog)[0][2].status=="待确认"


def test_foreign_proxy_gate(demo,project,catalog):
    row=dict(demo[5],确认代理="否")
    r=run([row],project,catalog)[0][0]
    assert r.status=="待确认" and "代理" in r.reason
    row["确认代理"]="是"
    r=run([row],project,catalog)[0][0]
    assert r.status=="已计算" and r.scope=="范围3"


def test_pending_factor_requires_confirmation(demo,project,catalog):
    catalog.factor("F031-v2023")["核对状态"]="待核对"
    assert run([demo[0]],project,catalog)[0][0].status=="待确认"
    r=run([dict(demo[0],确认待核因子="是")],project,catalog)[0][0]
    assert r.status=="已计算" and r.verification=="待核对"


def test_location_fallback(demo,project,catalog):
    row=dict(demo[0],地区="")
    r=run([row],Project(province=""),catalog)[0][0]
    assert r.emission_kg==pytest.approx(10000*.5306) and r.warnings
    catalog.factors=[f for f in catalog.factors if f["因子ID"]!="F031-v2023"]
    r=run([demo[0]],project,catalog)[0][0]
    assert r.emission_kg==pytest.approx(10000*.2472)
    catalog.factors=[f for f in catalog.factors if f["地区"]!="西南"]
    r=run([demo[0]],project,catalog)[0][0]
    assert r.emission_kg==pytest.approx(10000*.5306)


def test_future_factor_not_selected(demo,project,catalog):
    catalog.factor("F031-v2023")["数据年份"]=2030
    r=run([demo[0]],project,catalog)[0][0]
    assert r.factor_id!="F031-v2023"


def test_electricity_scenario_is_separate(demo,project,catalog):
    rs,_=run(demo,project,catalog)
    before=aggregate(rs)["total_kg"]
    scenario=electricity_scenario(rs,5000,catalog,project)
    assert scenario["情景kgCO2e"]==pytest.approx((10000+20*1000-5000)*.6096)
    assert aggregate(rs)["total_kg"]==before
    for green in [-1,30001]:
        with pytest.raises(UnitError):electricity_scenario(rs,green,catalog,project)
    assert electricity_scenario(rs,0,catalog,project)["情景kgCO2e"]==pytest.approx(30000*.6096)


def test_recompute_and_mutually_exclusive_groups(demo,project,catalog):
    rows=[dict(demo[0],计算组="同一活动"),dict(demo[0],计算组="同一活动")]
    rs,_=run(rows,project,catalog)
    assert aggregate(rs)["calculated_count"]==0
    overrides={rs[1].activity_id:{"不计入":"是","排除原因":"替代计量"}}
    first,audit=run(rows,project,catalog,overrides)
    second,audit2=run(rows,project,catalog,overrides)
    assert aggregate(first)["total_kg"]==pytest.approx(10000*.5581)
    assert aggregate(second)==aggregate(first) and audit.run_id==audit2.run_id


def test_manual_confirmation_rechecks_conditions(demo,project,catalog):
    rows=[demo[9]];rs,_=run(rows,project,catalog);aid=rs[0].activity_id
    override={aid:{"用途":"移动","密度kg/L":.84,"密度来源":"测试输入","确认适用条件":"是"}}
    r=run(rows,project,catalog,override)[0][0]
    assert r.emission_kg==pytest.approx(50*.84/1000*3186.3)
    override[aid]["确认因子ID"]="F031-v2023"
    assert run(rows,project,catalog,override)[0][0].status=="待确认"


def test_export_snapshot_and_formula_injection(demo,project,catalog):
    rs,audit=run(demo,project,catalog)
    rs[0].name="=HYPERLINK(\"https://example.com\",\"x\")"
    wb=openpyxl.load_workbook(BytesIO(export_excel(rs,project,catalog,audit)))
    assert {"汇总","计算明细","待处理清单","假设与口径","使用因子快照","使用GWP快照"}.issubset(wb.sheetnames)
    assert wb['计算明细']['B2'].data_type!="f"
    assert wb['待处理清单'].max_row==4
    assert wb['汇总']['B4'].value==pytest.approx(16351.19746)


def test_catalog_validation():
    w=openpyxl.Workbook();b=BytesIO();w.save(b)
    with pytest.raises(CatalogError):Catalog.load(b.getvalue())


def test_outside_inventory_gas_and_overflow(project,catalog):
    outside=dict(catalog.factor("GAS26-R410A"),排放表示="R22",GWP键="AR5-R22")
    with pytest.raises(UnitError,match="单列"):calculate(1,outside,catalog,project)
    huge=dict(catalog.factor("F031-v2023"),因子值=1e308)
    with pytest.raises(UnitError,match="范围"):calculate(1e308,huge,catalog,project)


def test_reporting_year_is_not_underlying_data_year(catalog):
    f=catalog.factor("UK26-17_404_4005_1_1")
    assert f["适用年度"]==2026 and f["数据年份"]==""
    assert catalog.factor("F031-v2023")["数据年份"]==2023
